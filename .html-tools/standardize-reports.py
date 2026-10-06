#!/usr/bin/env python3
"""
Standardize table styling across every CSW compliance report so the whole set
matches the gold-standard look (PCI DSS / NIST 800-53 / HIPAA):

  * dark-blue header row (#003366) with white bold text
  * single light borders on every table
  * zebra striping (white / #F2F2F2) on body rows
  * color-coded Coverage/status cells:
        Direct             -> #107C41 (green)
        Supporting         -> #C55A11 (amber)
        Evidence Required  -> #808080 (grey)   (customer supplies evidence elsewhere)
  * one shared vocabulary everywhere:
        "Full Coverage"    -> "Direct"
        "Partial Coverage" -> "Supporting"
        "Out of scope"     -> "Evidence Required"

Reports are classified automatically:
  * PLAIN  reports (no existing cell shading) get the FULL styler.
  * STYLED reports (already hand-formatted, the gold standard) get vocabulary
    normalization ONLY, so their bespoke layout and colors are preserved.

After styling, every report .docx is reconverted to .pdf with LibreOffice so the
PDF review copies pick up the new look. (HTML is rebuilt separately by
build-html.py.)

Usage:
    python3 .html-tools/standardize-reports.py
"""

from __future__ import annotations

import glob
import re
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

ROOT = Path(__file__).resolve().parent.parent

# Gold-standard palette (lifted from the existing styled reports).
HEADER_FILL = "003366"
ZEBRA_FILL = "F2F2F2"
WHITE = "FFFFFF"
BORDER = "BFBFBF"
FILL = {"direct": "107C41", "supporting": "C55A11", "evidence": "808080"}

# Vocabulary normalization for exact-match status cells.
VOCAB = {
    "full coverage": "Direct",
    "full": "Direct",
    "partial coverage": "Supporting",
    "partial": "Supporting",
    "out of scope": "Evidence Required",
    "out-of-scope": "Evidence Required",
}


def classify(text: str) -> str | None:
    """Return 'direct' | 'supporting' | 'evidence' for a short status cell."""
    t = text.strip().lower()
    if not t or len(t) > 30:
        return None
    # Evidence first (so "out of scope" / "evidence required" win).
    if re.match(r"^(evidence required|out of scope|out-of-scope|not applicable|n/a|none)\b", t):
        return "evidence"
    if re.match(r"^(direct|full coverage|full)\b", t):
        return "direct"
    if re.match(r"^(supporting|partial coverage|partial)\b", t):
        return "supporting"
    return None


def canonical(text: str, cls: str) -> str:
    """Canonical label, preserving a trailing parenthetical qualifier."""
    t = text.strip()
    if cls == "evidence":
        return "Evidence Required"
    m = re.search(r"(\([^)]*\))\s*$", t)
    qual = f" {m.group(1)}" if m else ""
    base = {"direct": "Direct", "supporting": "Supporting"}[cls]
    return base + qual


def set_cell_shading(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for shd in tcPr.findall(qn("w:shd")):
        tcPr.remove(shd)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def recolor_runs(cell, color: str, bold: bool = True) -> None:
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor.from_string(color)
            r.bold = bold


def set_cell_text(cell, text: str, color: str, bold: bool = True) -> None:
    """Replace cell content with a single styled run."""
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    run = p.add_run(text)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_text_keepfmt(cell, text: str) -> None:
    """Replace cell text but keep the first run's formatting (for styled reports)."""
    p = cell.paragraphs[0]
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def set_table_borders(tbl) -> None:
    tblPr = tbl._tbl.tblPr
    for b in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(b)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER)
        borders.append(el)
    tblPr.append(borders)


def style_table_full(tbl) -> None:
    nrows = len(tbl.rows)
    ncols = len(tbl.columns)
    set_table_borders(tbl)
    has_header = nrows >= 2 and ncols >= 2
    if has_header:
        for cell in tbl.rows[0].cells:
            set_cell_shading(cell, HEADER_FILL)
            recolor_runs(cell, WHITE, bold=True)
    start = 1 if has_header else 0
    for ri, row in enumerate(tbl.rows[start:], start=start):
        for cell in row.cells:
            cls = classify(cell.text)
            if cls:
                set_cell_shading(cell, FILL[cls])
                set_cell_text(cell, canonical(cell.text, cls), WHITE, bold=True)
            elif has_header and ri % 2 == 0:
                set_cell_shading(cell, ZEBRA_FILL)


def normalize_vocab(doc) -> int:
    """Exact-match vocabulary normalization that preserves existing formatting."""
    changed = 0
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                key = cell.text.strip().lower()
                if key in VOCAB and cell.text.strip() != VOCAB[key]:
                    set_text_keepfmt(cell, VOCAB[key])
                    changed += 1
    return changed


# Legend paragraphs (mid-migration prose) rewritten to the unified vocabulary.
# Each entry: (startswith-match, bold-lead, remaining-text). Empty lead => plain.
LEGEND_REWRITES = [
    (
        "Full Coverage on the older reports",
        "Direct",
        " means Secure Workload produces the primary evidence artifact for that row "
        "\u2014 segmentation, policy, flow and process telemetry, inventory, "
        "vulnerability and reachability, or forensic reconstruction. Direct is still "
        "evidence, not an attestation.",
    ),
    (
        "Partial Coverage means",
        "Supporting",
        " means Secure Workload contributes an input that feeds or complements another "
        "control, tool, or process \u2014 detection input only, reconciled against an "
        "external register, or one log source among several.",
    ),
    (
        "Evidence Required means Secure Workload has no record",
        "Evidence Required",
        " means Secure Workload has no record for that row. The customer supplies the "
        "contract, the analysis, or the HR / physical-security evidence from another "
        "control.",
    ),
    (
        "Direct, Supporting, and Out of scope on the newer reports",
        "",
        "Every report in this set now uses the same three coverage words \u2014 Direct, "
        "Supporting, and Evidence Required. Coverage describes what evidence Secure "
        "Workload produces; it never means the control is passed.",
    ),
]


def set_para_text(p, text: str) -> None:
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.add_run(text)


def rewrite_para(p, lead: str, rest: str) -> None:
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    r1 = p.add_run(lead)
    r1.bold = True
    p.add_run(rest)


def normalize_prose(doc) -> int:
    """Unify legend + inline coverage vocabulary in body paragraphs."""
    changed = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("How to read the coverage words"):
            set_para_text(p, "How to read the Coverage column")
            changed += 1
            continue
        for start, lead, rest in LEGEND_REWRITES:
            if t.startswith(start):
                if lead:
                    rewrite_para(p, lead, rest)
                else:
                    set_para_text(p, rest)
                changed += 1
                break
    # Inline term normalization for remaining analytical prose / lead-in labels.
    for p in doc.paragraphs:
        for r in p.runs:
            if "Full Coverage" in r.text or "Partial Coverage" in r.text:
                r.text = (
                    r.text.replace("Full Coverage", "Direct coverage")
                    .replace("Partial Coverage", "Supporting coverage")
                )
                changed += 1
    return changed


def is_styled(docx_path: Path) -> bool:
    """A report is 'already styled' if its document.xml carries cell shading."""
    try:
        with zipfile.ZipFile(docx_path) as z:
            xml = z.read("word/document.xml").decode("utf-8", "replace")
    except Exception:
        return False
    return "w:shd " in xml or "w:shd>" in xml


def style_report(docx_path: Path) -> str:
    doc = Document(str(docx_path))
    prose = normalize_prose(doc)
    if is_styled(docx_path):
        n = normalize_vocab(doc)
        doc.save(str(docx_path))
        return f"vocab-only ({n} cell + {prose} prose edits)"
    for tbl in doc.tables:
        style_table_full(tbl)
    doc.save(str(docx_path))
    return f"full style ({len(doc.tables)} table(s), {prose} prose edits)"


def convert_pdf(docx_path: Path) -> bool:
    out_dir = docx_path.parent
    cmd = [
        "soffice",
        "--headless",
        "-env:UserInstallation=file:///tmp/lo_profile_csw",
        "--convert-to",
        "pdf",
        "--outdir",
        str(out_dir),
        str(docx_path),
    ]
    try:
        subprocess.run(cmd, check=True, capture_output=True, text=True, timeout=180)
        return True
    except Exception as e:
        print(f"    PDF FAILED: {e}")
        return False


def main() -> int:
    reports = sorted(ROOT.glob("*/CSW-*-Compliance-Report.docx"))
    if not reports:
        print("No report docx found.", file=sys.stderr)
        return 1
    print(f"Styling {len(reports)} report(s)...\n")
    for docx_path in reports:
        rel = docx_path.relative_to(ROOT)
        mode = style_report(docx_path)
        ok = convert_pdf(docx_path)
        pdf = "pdf OK" if ok else "pdf ERR"
        print(f"  {rel}  ->  {mode}; {pdf}")
    print("\nDone. Rebuild HTML with .html-tools/build-html.py")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
